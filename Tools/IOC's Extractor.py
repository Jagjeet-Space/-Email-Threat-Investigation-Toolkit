#!/usr/bin/env python3
import re
import sys
import json
import os
import email
from email import policy
from email.parser import BytesParser
from typing import Dict, Set, Any
import tldextract
import magic
import hashlib
from PyPDF2 import PdfReader
import docx  # <-- NEW: Library for Word files
import openpyxl  # <-- NEW: Library for Excel files

# This is your existing function. It remains unchanged.
def find_iocs(text: str) -> Dict[str, Set[str]]:
    """
    Finds and returns a dictionary of unique IOCs from a given string.
    """
    broad_domain_pattern = r'\b(?:[a-zA-Z0-9-]+\.){1,}[a-zA-Z]{2,}\b'
    ip_pattern = r'\b(?:[0-9]{1,3}\.){3}[0-9]{1,3}\b'
    url_pattern = r'https?://[^\s<>"]+'
    
    iocs = {
        "ips": set(re.findall(ip_pattern, text)),
        "urls": set(re.findall(url_pattern, text)),
        "domains": set()
    }

    domain_candidates = re.findall(broad_domain_pattern, text)

    for candidate in domain_candidates:
        extracted = tldextract.extract(candidate)
        if extracted.domain and extracted.suffix:
            full_domain = f"{extracted.domain}.{extracted.suffix}"
            iocs["domains"].add(full_domain)
            
    return iocs

# --- MODIFIED: This function now supports more file types. ---
def parse_file_content(file_path: str) -> str:
    """Parses a file for text content based on its type."""
    try:
        mime_type = magic.from_file(file_path, mime=True)
        print(f"[*] Detected file type: {mime_type}", file=sys.stderr)
        
        if 'text/' in mime_type:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        
        elif 'application/pdf' in mime_type:
            content = ""
            with open(file_path, 'rb') as f:
                reader = PdfReader(f)
                for page in reader.pages:
                    content += page.extract_text() or ""
            return content
        
        # NEW: Word Document Parsing
        elif 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' in mime_type:
            doc = docx.Document(file_path)
            content = " ".join([para.text for para in doc.paragraphs])
            return content
        
        # NEW: Excel Spreadsheet Parsing
        elif 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in mime_type:
            workbook = openpyxl.load_workbook(file_path)
            content = ""
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value:
                            content += str(cell.value) + " "
            return content
            
        else:
            return ""
            
    except Exception as e:
        print(f"[!] Error parsing file {file_path}: {e}", file=sys.stderr)
        return ""

# --- MODIFIED: This function now includes header parsing and better file handling. ---
def get_iocs_from_file_or_content(input_data, is_file_path=True) -> Dict[str, Any]:
    """
    Primary function to get IOCs from a file path or raw content,
    now with header analysis and expanded file support.
    """
    attributed_iocs = {
        "headers": {}, # <-- NEW: Placeholder for header IOCs
        "body": { "ips": set(), "urls": set(), "domains": set() },
        "attachments": {}
    }

    if is_file_path and os.path.exists(input_data) and (input_data.endswith('.eml') or input_data.endswith('.msg')):
        print(f"[*] Processing email file: {input_data}", file=sys.stderr)
        with open(input_data, 'rb') as fp:
            msg = BytesParser(policy=policy.default).parse(fp)
            
        # --- NEW LOGIC: EXTRACT AND ANALYZE HEADERS ---
        header_text = ""
        for key, value in msg.items():
            header_text += f"{key}: {value}\n"
        
        header_iocs = find_iocs(header_text)
        attributed_iocs["headers"] = {key: list(val) for key, val in header_iocs.items()}

        # ... (Body parsing logic is the same)
        email_body_text = ""
        if msg.is_multipart():
            for part in msg.walk():
                content_type = part.get_content_type()
                if content_type == 'text/plain' or content_type == 'text/html':
                    email_body_text += part.get_content() or ""
        else:
            email_body_text = msg.get_content() or ""
        
        body_iocs = find_iocs(email_body_text)
        attributed_iocs["body"] = {key: list(val) for key, val in body_iocs.items()}
            
        # --- MODIFIED ATTACHMENT LOGIC: Better file type handling ---
        for part in msg.walk():
            if part.get_content_disposition() == 'attachment':
                filename = part.get_filename() or "unnamed_attachment"
                temp_filename = f"temp_{filename}"
                with open(temp_filename, 'wb') as f:
                    f.write(part.get_payload(decode=True))

                print(f"[*] Analyzing attachment: {filename}", file=sys.stderr)
                
                # Get hash of the attachment
                with open(temp_filename, 'rb') as f:
                    data = f.read()
                    md5_hash = hashlib.md5(data).hexdigest()
                    sha256_hash = hashlib.sha256(data).hexdigest()
                
                # Check file type for content parsing
                mime_type = magic.from_file(temp_filename, mime=True)
                
                attachment_content = ""
                # Only try to parse content if it's not a binary executable
                if not any(ext in mime_type for ext in ['application/x-msdownload', 'application/octet-stream']):
                    attachment_content = parse_file_content(temp_filename)
                
                attachment_iocs = find_iocs(attachment_content)
                
                attributed_iocs["attachments"][filename] = {
                    "hashes": {
                        "md5": md5_hash,
                        "sha256": sha256_hash
                    },
                    "iocs": {key: list(val) for key, val in attachment_iocs.items()}
                }
                
                os.remove(temp_filename)
    
    else:
        text_content = input_data if not is_file_path else parse_file_content(input_data)
        if text_content:
            found_iocs = find_iocs(text_content)
            attributed_iocs["body"] = {key: list(val) for key, val in found_iocs.items()}

    return attributed_iocs

# --- MODIFIED: Final output block now handles the new "headers" key. ---
if __name__ == "__main__":
    content = ""
    attributed_iocs = {}
    
    if len(sys.argv) > 1:
        input_path = sys.argv[1]
        if os.path.exists(input_path) and os.path.isfile(input_path):
            attributed_iocs = get_iocs_from_file_or_content(input_path, is_file_path=True)
        else:
            print(f"Error: File not found or is not a file: {input_path}", file=sys.stderr)
            sys.exit(1)
            
    elif not sys.stdin.isatty():
        content = sys.stdin.read()
        attributed_iocs = get_iocs_from_file_or_content(content, is_file_path=False)
        
    else:
        print("Please provide a file path as an argument or pipe content to standard input.", file=sys.stderr)
        sys.exit(1)

    json_output_data = {}
    
    # Process headers
    if attributed_iocs.get("headers"):
        cleaned_headers_iocs = {k: sorted(list(v)) for k, v in attributed_iocs["headers"].items() if v}
        if cleaned_headers_iocs:
            json_output_data["headers"] = cleaned_headers_iocs
    
    # Process body
    if attributed_iocs.get("body"):
        cleaned_body_iocs = {k: sorted(list(v)) for k, v in attributed_iocs["body"].items() if v}
        if cleaned_body_iocs:
            json_output_data["body"] = cleaned_body_iocs
    
    # Process attachments
    if attributed_iocs.get("attachments"):
        cleaned_attachments = {}
        for filename, attachment_data in attributed_iocs["attachments"].items():
            cleaned_iocs = {k: sorted(list(v)) for k, v in attachment_data["iocs"].items() if v}
            if any(cleaned_iocs.values()) or any(attachment_data["hashes"].values()):
                 cleaned_attachments[filename] = {
                    "hashes": attachment_data["hashes"],
                    "iocs": cleaned_iocs
                 }
        if cleaned_attachments:
            json_output_data["attachments"] = cleaned_attachments

    final_output = {
        "IOCs Extracted": json_output_data
    }

    if not final_output["IOCs Extracted"]:
        print(json.dumps({"info": "No IOCs found."}), indent=2)
    else:
        print(json.dumps(final_output, indent=2))