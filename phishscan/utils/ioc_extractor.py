#!/usr/bin/env python3
import re
import sys
import json
import os
from typing import Dict, Set, Any, Tuple, Optional

# Email parsing
from email import policy
from email.parser import BytesParser

# Optional colorized output
try:
    from termcolor import colored as _colored
except Exception:
    def _colored(text, color=None):
        return text

def make_colorer(disable: bool = False):
    if disable:
        return lambda text, color=None: text
    return _colored

# Optional libs (graceful fallbacks)
# tldextract
try:
    import tldextract
except Exception:
    tldextract = None

# libmagic (file type detection)
try:
    import magic  # python-magic (linux/mac) or python-magic-bin (windows)
except Exception:
    magic = None

# PDF
try:
    from PyPDF2 import PdfReader
except Exception:
    PdfReader = None

# Word
try:
    import docx
except Exception:
    docx = None

# Excel
try:
    import openpyxl
except Exception:
    openpyxl = None

import hashlib


# ---------------------------
# IOC extraction primitives
# ---------------------------

def find_iocs(text: str) -> Dict[str, Set[str]]:
    """
    Finds and returns a dictionary of unique IOCs from a given string.
    Returns sets for internal use; callers should convert to lists before JSON.
    """
    # Broad domain candidate pattern (will normalize via tldextract)
    broad_domain_pattern = r'\b(?:[a-zA-Z0-9-]+\.){1,}[a-zA-Z]{2,}\b'
    ip_pattern = r'\b(?:[0-9]{1,3}\.){3}[0-9]{1,3}\b'
    url_pattern = r'https?://[^\s<>"\']+'

    iocs = {
        "ips": set(re.findall(ip_pattern, text or "")),
        "urls": set(re.findall(url_pattern, text or "")),
        "domains": set()
    }

    domain_candidates = re.findall(broad_domain_pattern, text or "")
    if tldextract is not None:
        for candidate in domain_candidates:
            try:
                ext = tldextract.extract(candidate)
                if ext.domain and ext.suffix:
                    full_domain = f"{ext.domain}.{ext.suffix}".lower()
                    iocs["domains"].add(full_domain)
            except Exception:
                iocs["domains"].add(candidate.lower())
    else:
        for candidate in domain_candidates:
            iocs["domains"].add(candidate.lower())

    return iocs


# ---------------------------
# File content parsers
# ---------------------------

def _detect_mime(file_path: str) -> Optional[str]:
    if magic is None:
        # Fallback to naive guess by extension when libmagic is not available
        ext = os.path.splitext(file_path)[1].lower()
        if ext in {".txt", ".log", ".csv", ".md", ".eml"}:
            return "text/plain"
        if ext == ".pdf":
            return "application/pdf"
        if ext == ".docx":
            return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        if ext == ".xlsx":
            return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        return "application/octet-stream"
    try:
        return magic.from_file(file_path, mime=True)
    except Exception:
        return None


def parse_file_content(file_path: str, quiet: bool = False) -> str:
    """
    Parses a file for text content based on its type.
    Supports text, PDF, DOCX, XLSX with graceful degradation if libs are missing.
    """
    try:
        mime_type = _detect_mime(file_path) or "application/octet-stream"
        if not quiet:
            print(f"[*] Detected file type: {mime_type}", file=sys.stderr)

        # Plain text
        if mime_type.startswith("text/"):
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        # PDF
        if "application/pdf" in mime_type:
            if PdfReader is None:
                if not quiet:
                    print("[!] PyPDF2 not available; skipping PDF text extraction", file=sys.stderr)
                return ""
            content = ""
            with open(file_path, "rb") as f:
                reader = PdfReader(f)
                for page in reader.pages:
                    content += (page.extract_text() or "")
            return content

        # Word (docx)
        if "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in mime_type:
            if docx is None:
                if not quiet:
                    print("[!] python-docx not available; skipping DOCX text extraction", file=sys.stderr)
                return ""
            d = docx.Document(file_path)
            return " ".join([para.text for para in d.paragraphs])

        # Excel (xlsx)
        if "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in mime_type:
            if openpyxl is None:
                if not quiet:
                    print("[!] openpyxl not available; skipping XLSX text extraction", file=sys.stderr)
                return ""
            wb = openpyxl.load_workbook(file_path, data_only=True)
            content = ""
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            content += f"{cell.value} "
            return content

        # Unknown/binary
        return ""

    except Exception as e:
        if not quiet:
            print(f"[!] Error parsing file {file_path}: {e}", file=sys.stderr)
        return ""


# ---------------------------
# Email parsing and IOC attribution
# ---------------------------

def _email_meta(msg) -> Dict[str, Any]:
    return {
        "subject": msg.get("Subject"),
        "from": msg.get("From"),
        "to": msg.get("To"),
        "date": msg.get("Date"),
        "message_id": msg.get("Message-ID") or msg.get("Message-Id"),
    }


def get_iocs_from_file_or_content(input_data, is_file_path=True, quiet: bool = False) -> Tuple[Dict[str, Any], Optional[Dict[str, Any]]]:
    """
    Primary function to get IOCs from a file path or raw content,
    with header analysis and expanded file support.

    Returns:
        (attributed_iocs: dict, meta: dict|None)
    """
    attributed_iocs: Dict[str, Any] = {
        "headers": {},
        "body": {"ips": set(), "urls": set(), "domains": set()},
        "attachments": {}
    }
    meta: Optional[Dict[str, Any]] = None

    # Email file path flow
    if is_file_path and os.path.exists(input_data) and (input_data.lower().endswith(".eml") or input_data.lower().endswith(".msg")):
        if not quiet:
            print(f"[*] Processing email file: {input_data}", file=sys.stderr)
        with open(input_data, "rb") as fp:
            msg = BytesParser(policy=policy.default).parse(fp)

        # Capture minimal metadata for banner
        meta = _email_meta(msg)

        # Headers IOC extraction
        header_text = ""
        for key, value in msg.items():
            header_text += f"{key}: {value}\n"
        header_iocs = find_iocs(header_text)
        attributed_iocs["headers"] = {k: sorted(v) for k, v in header_iocs.items()}

        # Body IOC extraction
        email_body_text = ""
        if msg.is_multipart():
            for part in msg.walk():
                ctype = part.get_content_type()
                if ctype in ("text/plain", "text/html"):
                    try:
                        email_body_text += part.get_content() or ""
                    except Exception:
                        pass
        else:
            try:
                email_body_text = msg.get_content() or ""
            except Exception:
                email_body_text = ""

        body_iocs = find_iocs(email_body_text)
        attributed_iocs["body"] = {k: sorted(v) for k, v in body_iocs.items()}

        # Attachments IOC extraction
        for part in msg.walk():
            if part.get_content_disposition() == "attachment":
                filename = part.get_filename() or "unnamed_attachment"
                safe_name = filename.replace(os.sep, "_")
                temp_filename = f"temp_{safe_name}"
                try:
                    with open(temp_filename, "wb") as f:
                        f.write(part.get_payload(decode=True) or b"")
                    if not quiet:
                        print(f"[*] Analyzing attachment: {filename}", file=sys.stderr)

                    # Hashes
                    with open(temp_filename, "rb") as f:
                        data = f.read()
                        md5_hash = hashlib.md5(data).hexdigest()
                        sha256_hash = hashlib.sha256(data).hexdigest()

                    # File type
                    mime_type = _detect_mime(temp_filename) or "application/octet-stream"

                    # Extract content if likely text-like
                    attachment_content = ""
                    likely_binary = any(sig in (mime_type or "") for sig in ["application/x-msdownload"])
                    if not likely_binary:
                        attachment_content = parse_file_content(temp_filename, quiet=quiet)

                    attachment_iocs = find_iocs(attachment_content)

                    attributed_iocs["attachments"][filename] = {
                        "hashes": {"md5": md5_hash, "sha256": sha256_hash},
                        "mimetype": mime_type,
                        "iocs": {k: sorted(v) for k, v in attachment_iocs.items()}
                    }
                finally:
                    try:
                        if os.path.exists(temp_filename):
                            os.remove(temp_filename)
                    except Exception:
                        pass

    else:
        # Non-email path or raw content flow
        if is_file_path and os.path.exists(input_data):
            text_content = parse_file_content(input_data, quiet=quiet)
        else:
            text_content = input_data if not is_file_path else ""
        if text_content:
            found_iocs = find_iocs(text_content)
            attributed_iocs["body"] = {k: sorted(v) for k, v in found_iocs.items()}

    return attributed_iocs, meta


# ---------------------------
# Pretty printer (polished)
# ---------------------------

def _plural(n, word):
    return f"{n} {word}{'' if n == 1 else 's'}"


def _section(title, colorer):
    print(colorer(title, "cyan"))


def _print_list(label, items, indent="  "):
    items = sorted(items or [])
    if not items:
        print(f"{indent}{label}: -")
        return
    print(f"{indent}{label}:")
    for it in items:
        print(f"{indent}  - {it}")


def print_iocs_pretty(attributed_iocs: Dict[str, Any], meta: Optional[Dict[str, Any]], colorer):
    # Banner
    print(colorer("[*] ioc_analyzer v1.0", "cyan"))
   # if meta:
   #     subj = meta.get("subject") or "-"
   #     frm = meta.get("from") or "-"
   #     to = meta.get("to") or "-"
   #     dt = meta.get("date") or "-"
   #     print(f"Subject: {subj}")
   #     print(f"From:    {frm}")
   #     print(f"To:      {to}")
   #     print(f"Date:    {dt}")
   #     print("")

    # Headers IOCs
    headers = attributed_iocs.get("headers", {}) or {}
    h_ips = headers.get("ips", []) or []
    h_urls = headers.get("urls", []) or []
    h_domains = headers.get("domains", []) or []
    _section("[*] Header IOC summary", colorer)
    print(f"  {_plural(len(h_ips), 'IP')}, {_plural(len(h_urls), 'URL')}, {_plural(len(h_domains), 'domain')}")
    _print_list("IPs", h_ips)
    _print_list("URLs", h_urls)
    _print_list("Domains", h_domains)
    print("")

    # Body IOCs
    body = attributed_iocs.get("body", {}) or {}
    b_ips = body.get("ips", []) or []
    b_urls = body.get("urls", []) or []
    b_domains = body.get("domains", []) or []
    _section("[*] Body IOC summary", colorer)
    print(f"  {_plural(len(b_ips), 'IP')}, {_plural(len(b_urls), 'URL')}, {_plural(len(b_domains), 'domain')}")
    _print_list("IPs", b_ips)
    _print_list("URLs", b_urls)
    _print_list("Domains", b_domains)
    print("")

    # Attachments
    atts = attributed_iocs.get("attachments", {}) or {}
    _section("[*] Attachments", colorer)
    if not atts:
        print("  None")
        return

    for fname, info in atts.items():
        print(f"- {fname}")
        hashes = (info.get("hashes") or {})
        md5 = hashes.get("md5") or "-"
        sha256 = hashes.get("sha256") or "-"
        mimetype = info.get("mimetype") or "-"
        print(f"    Type:   {mimetype}")
        print(f"    MD5:    {md5}")
        print(f"    SHA256: {sha256}")

        iocs = info.get("iocs", {}) or {}
        a_ips = iocs.get("ips", []) or []
        a_urls = iocs.get("urls", []) or []
        a_domains = iocs.get("domains", []) or []
        print(f"    IOC counts: {_plural(len(a_ips), 'IP')}, {_plural(len(a_urls), 'URL')}, {_plural(len(a_domains), 'domain')}")
        if a_ips:
            _print_list("IPs", a_ips, indent="    ")
        if a_urls:
            _print_list("URLs", a_urls, indent="    ")
        if a_domains:
            _print_list("Domains", a_domains, indent="    ")
        print("")


# ---------------------------
# CLI
# ---------------------------

def main():
    import argparse
    ap = argparse.ArgumentParser(description="Extract and display IOCs from files or .eml emails (pretty console output).")
    ap.add_argument("input", help="Path to a file (e.g., .eml, .txt, .pdf, .docx, .xlsx) or raw text if --raw is set")
    ap.add_argument("--raw", action="store_true", help="Treat input as raw text instead of a file path")
    ap.add_argument("--json-out", help="Write the IOC result to a JSON file")
    ap.add_argument("--print-json", action="store_true", help="Also print the JSON to console (in addition to pretty output)")
    ap.add_argument("--format", choices=["pretty", "json", "csv"], default="pretty", help="Output format (default: pretty)")
    ap.add_argument("--quiet", "-q", action="store_true", help="Suppress stderr diagnostics")
    ap.add_argument("--no-color", action="store_true", help="Disable colored terminal output")
    args = ap.parse_args()

    colorer = make_colorer(disable=args.no_color)

    if args.raw:
        result, meta = get_iocs_from_file_or_content(args.input, is_file_path=False, quiet=args.quiet)
    else:
        if not os.path.exists(args.input):
            print(f"[!] File not found: {args.input}", file=sys.stderr)
            sys.exit(1)
        result, meta = get_iocs_from_file_or_content(args.input, is_file_path=True, quiet=args.quiet)

    # Display format
    if args.format == "pretty":
        print_iocs_pretty(result, meta, colorer)
    elif args.format == "json":
        print(json.dumps(result, indent=2, ensure_ascii=False))
    elif args.format == "csv":
        import csv
        rows = []
        # headers
        for x in result.get("headers", {}).get("ips", []): rows.append(("headers", "ip", x))
        for x in result.get("headers", {}).get("urls", []): rows.append(("headers", "url", x))
        for x in result.get("headers", {}).get("domains", []): rows.append(("headers", "domain", x))
        # body
        for x in result.get("body", {}).get("ips", []): rows.append(("body", "ip", x))
        for x in result.get("body", {}).get("urls", []): rows.append(("body", "url", x))
        for x in result.get("body", {}).get("domains", []): rows.append(("body", "domain", x))
        # attachments
        for fname, info in (result.get("attachments") or {}).items():
            i = (info.get("iocs") or {})
            for x in i.get("ips", []): rows.append((f"attachment:{fname}", "ip", x))
            for x in i.get("urls", []): rows.append((f"attachment:{fname}", "url", x))
            for x in i.get("domains", []): rows.append((f"attachment:{fname}", "domain", x))
        w = csv.writer(sys.stdout)
        w.writerow(["section", "type", "indicator"])
        w.writerows(rows)

    # Optional JSON file output regardless of display format
    if args.json_out:
        try:
            with open(args.json_out, "w", encoding="utf-8") as f:
                json.dump(result, f, ensure_ascii=False, indent=2)
            if not args.quiet:
                print(colorer(f"[*] Wrote JSON to {args.json_out}", "green"), file=sys.stderr)
        except Exception as e:
            print(f"[!] Failed to write JSON: {e}", file=sys.stderr)
            sys.exit(2)

    # Optional: also print JSON to console in addition to chosen format
    if args.print_json and args.format != "json":
        print(json.dumps(result, indent=2, ensure_ascii=False))

    sys.exit(0)


if __name__ == "__main__":
    main()
